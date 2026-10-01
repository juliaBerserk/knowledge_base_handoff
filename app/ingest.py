from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.secrets import redact

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Формат {suffix} не поддерживается. Используйте PDF, DOCX, TXT или MD.")
    if suffix in {".txt", ".md"}:
        raw = path.read_text(encoding="utf-8", errors="ignore")
    elif suffix == ".pdf":
        reader = PdfReader(str(path))
        raw = "\n".join((page.extract_text() or "") for page in reader.pages)
    else:
        doc = Document(str(path))
        parts: list[str] = []
        for para in doc.paragraphs:
            parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        raw = "\n".join(parts)
    text = redact(raw)
    text = "\n".join(line.rstrip() for line in text.splitlines())
    text = "\n\n".join(block.strip() for block in text.split("\n\n") if block.strip())
    if not text.strip():
        raise ValueError("Не удалось извлечь текст из файла.")
    return text


def chunk_text(text: str, size: int, overlap: int) -> list[str]:
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    chunks: list[str] = []
    buf = ""
    for para in paragraphs:
        candidate = f"{buf}\n{para}".strip() if buf else para
        if len(candidate) <= size:
            buf = candidate
            continue
        if buf:
            chunks.append(buf)
        if len(para) <= size:
            buf = para
        else:
            start = 0
            while start < len(para):
                end = start + size
                chunks.append(para[start:end])
                start = max(end - overlap, start + 1)
            buf = ""
    if buf:
        chunks.append(buf)

    merged: list[str] = []
    for chunk in chunks:
        if merged and len(merged[-1]) < size // 3:
            merged[-1] = f"{merged[-1]}\n{chunk}"
        else:
            merged.append(chunk)
    return merged
